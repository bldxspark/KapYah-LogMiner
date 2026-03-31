"""KapYah LogMiner report export pipeline.

Maintained by Durgesh Tiwari, KapYah Industries Pvt. Ltd.

File purpose: Python report generator that writes the Excel workbook and PDF mission report.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
import sys
from typing import Any
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MAX_RC_CHANNELS = 16
RC_CHANNEL_COLUMNS = [f"rc{index}" for index in range(1, MAX_RC_CHANNELS + 1)]


def _apply_series_colors(chart: LineChart | BarChart, colors: list[str]) -> None:
    for index, color in enumerate(colors):
        if index >= len(chart.ser):
            break
        chart.ser[index].graphicalProperties.line.solidFill = color
        chart.ser[index].graphicalProperties.line.width = 28000
        if hasattr(chart.ser[index].graphicalProperties, "solidFill"):
            chart.ser[index].graphicalProperties.solidFill = color


def _autosize_columns(sheet: Worksheet) -> None:
    for column_cells in sheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 36)


def _add_line_chart(
    sheet: Worksheet,
    *,
    title: str,
    y_axis_title: str,
    x_axis_title: str,
    min_col: int,
    max_col: int,
    header_row: int,
    first_data_row: int,
    last_data_row: int,
    anchor: str,
    colors: list[str],
    height: float = 9,
    width: float = 18,
) -> None:
    if last_data_row - first_data_row + 1 < 2:
        return

    chart = LineChart()
    chart.title = title
    chart.style = 2
    chart.y_axis.title = y_axis_title
    chart.x_axis.title = x_axis_title
    chart.height = height
    chart.width = width
    chart.add_data(
        Reference(sheet, min_col=min_col, max_col=max_col, min_row=header_row, max_row=last_data_row),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=first_data_row, max_row=last_data_row))
    _apply_series_colors(chart, colors)
    sheet.add_chart(chart, anchor)


def _add_bar_chart(
    sheet: Worksheet,
    *,
    title: str,
    y_axis_title: str,
    min_col: int,
    max_col: int,
    header_row: int,
    first_data_row: int,
    last_data_row: int,
    anchor: str,
    colors: list[str],
) -> None:
    if last_data_row - first_data_row + 1 < 1:
        return

    chart = BarChart()
    chart.title = title
    chart.style = 10
    chart.y_axis.title = y_axis_title
    chart.height = 7
    chart.width = 12
    chart.add_data(
        Reference(sheet, min_col=min_col, max_col=max_col, min_row=header_row, max_row=last_data_row),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=first_data_row, max_row=last_data_row))
    _apply_series_colors(chart, colors)
    sheet.add_chart(chart, anchor)


def _normalize_requested_output_path(output_dir: str) -> Path:
    # Treat the dialog result as a folder name even if a suffix is typed accidentally.
    cleaned = output_dir.strip().strip('"').strip("'")
    requested_path = Path(cleaned).expanduser()
    if requested_path.suffix:
        requested_path = requested_path.with_suffix("")
    return Path(str(requested_path))


def _next_available_folder(report_root: Path, requested_name: str) -> Path:
    # Preserve older exports by choosing the next free folder name when needed.
    base_name = requested_name.strip().rstrip(". ") or "report"
    report_folder = report_root / base_name
    if not report_folder.exists():
        return report_folder

    match = re.match(r"^(.*?)(?:_(\d+))?$", base_name)
    if match:
        prefix = (match.group(1) or "report").rstrip("_") or "report"
        start_number = int(match.group(2)) + 1 if match.group(2) else 2
    else:
        prefix = base_name
        start_number = 2

    candidate_number = start_number
    while True:
        candidate = report_root / f"{prefix}_{candidate_number}"
        if not candidate.exists():
            return candidate
        candidate_number += 1



def _create_report_folder(output_dir: str | None) -> Path:
    if output_dir:
        requested_path = _normalize_requested_output_path(output_dir)
        report_root = requested_path.parent
        report_root.mkdir(parents=True, exist_ok=True)
        report_folder = _next_available_folder(report_root, requested_path.name)
        report_folder.mkdir(parents=True, exist_ok=False)
        return report_folder

    report_root = Path.home() / "Downloads"
    report_root.mkdir(parents=True, exist_ok=True)
    existing_reports = [path for path in report_root.glob("report_*") if path.is_dir()]
    report_folder = report_root / f"report_{len(existing_reports) + 1}"
    report_folder.mkdir(exist_ok=True)
    return report_folder


def _text(value: Any, fallback: str = "Unavailable") -> str:
    if value is None:
        return fallback
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or fallback
    if isinstance(value, list):
        cleaned = [str(item).strip() for item in value if str(item).strip()]
        return ", ".join(cleaned) if cleaned else fallback
    return str(value)


def _time_window_sentence(window: str | None) -> str:
    if not window or " - " not in window:
        return "The exact mission start and end window could not be derived from the reviewed log."
    start_time, end_time = [part.strip() for part in window.split(" - ", 1)]
    return f"The detected mission window runs from {start_time} to {end_time}."


def _assessment_label(overview: dict[str, Any], messages: dict[str, Any]) -> str:
    if overview.get("errorMessages") or overview.get("failsafeEvents") or (messages.get("errorCount") or 0) > 0:
        return "Requires Investigation"
    if overview.get("keyWarnings") or overview.get("keyAnomalies") or (messages.get("warningCount") or 0) > 0:
        return "Caution"
    return "Normal"


def _format_time(value: Any) -> str:
    if value is None:
        return "Unavailable"
    if isinstance(value, (int, float)):
        total_seconds = max(int(value), 0)
        minutes, seconds = divmod(total_seconds, 60)
        hours, minutes = divmod(minutes, 60)
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    return str(value)


def _join_limited(values: list[str], fallback: str = "None recorded", limit: int = 5) -> str:
    cleaned = [value.strip() for value in values if value and value.strip()]
    if not cleaned:
        return fallback
    return ", ".join(cleaned[:limit])


def _bundled_asset_root() -> Path | None:
    bundled_root = getattr(sys, "_MEIPASS", None)
    return Path(bundled_root) if bundled_root else None


def _resolve_logo_path() -> Path | None:
    bundled_asset_root = _bundled_asset_root()
    if bundled_asset_root:
        bundled_logo = bundled_asset_root / "assets" / "kapyah-company-mark-redico.png"
        if bundled_logo.exists():
            return bundled_logo

    workspace_root = Path(__file__).resolve().parents[1]
    workspace_logo = workspace_root / "src" / "assets" / "kapyah-company-mark-redico.png"
    if workspace_logo.exists():
        return workspace_logo

    return None


def _build_mission_overview_paragraphs(analysis: dict[str, Any], generated_at: datetime) -> list[str]:
    overview = analysis["overview"]
    timeline = analysis["timeline"]
    messages = analysis["messages"]
    report_time = generated_at.strftime("%d %b %Y at %I:%M %p")
    paragraph_one = (
        f"This KapYah mission analysis report was generated on {report_time} using the local KapYah LogMiner workflow. "
        f"The reviewed file {_text(overview.get('logName'))} is identified as a {_text(overview.get('vehicleType')).lower()} mission log. "
        f"The total detected flight duration is {_text(overview.get('totalFlightDuration'))}, and {_time_window_sentence(overview.get('armDisarmTime'))}"
    )
    paragraph_two = (
        f"The analysis pipeline reviewed {timeline.get('totalEvents', 0)} timeline events and {messages.get('totalMessages', 0)} captured messages. "
        f"Mission status is currently assessed as {_assessment_label(overview, messages)} based on the detected warnings, anomalies, failsafe indicators, and message severity counts."
    )
    paragraph_three = (
        f"Primary flight mode activity includes {_join_limited(overview.get('flightModes', []), fallback='no confirmed modes')}. "
        f"The mission recorded {_text(overview.get('flightCount'), fallback='0')} detected flight segment(s), with orientation sourced from {_text(overview.get('orientationSource'))}."
    )
    return [paragraph_one, paragraph_two, paragraph_three]


def _build_navigation_paragraphs(analysis: dict[str, Any]) -> list[str]:
    overview = analysis["overview"]
    map_data = analysis["map"]
    route_points = map_data.get("routePoints", [])
    paragraph_one = (
        f"GPS review indicates {_text(overview.get('gpsStatus'))} with {_text(overview.get('satelliteCount'))} satellites observed near the home position. "
        f"The mission home location is reported as {_text(overview.get('homeLocation'))}, and the estimated distance traveled is {_text(overview.get('distanceTraveled'))}."
    )
    paragraph_two = (
        f"A total of {len(route_points)} route samples were available for playback and map reconstruction. "
        f"Maximum height reached was {_text(overview.get('maxAltitude'))}, while maximum speed peaked at {_text(overview.get('maxSpeed'))}."
    )
    paragraph_three = (
        "Map playback in the desktop application remains synchronized with the reconstructed route, enabling event review alongside route progress, "
        "drone-follow behavior, and resettable mission perspective controls."
    )
    return [paragraph_one, paragraph_two, paragraph_three]


def _build_timeline_message_paragraphs(analysis: dict[str, Any]) -> list[str]:
    timeline = analysis["timeline"]
    messages = analysis["messages"]
    overview = analysis["overview"]
    paragraph_one = (
        f"Timeline review captured {timeline.get('totalEvents', 0)} events across the mission window. "
        f"Synthetic mission boundary markers are included so operators can quickly identify start and end conditions even when the raw log is incomplete."
    )
    paragraph_two = (
        f"Message review counted {messages.get('infoCount', 0)} informational, {messages.get('warningCount', 0)} warning, and {messages.get('errorCount', 0)} error messages. "
        f"Key warnings include {_join_limited(overview.get('keyWarnings', []))}, while notable anomalies include {_join_limited(overview.get('keyAnomalies', []))}."
    )
    paragraph_three = (
        f"Failsafe-related observations for this mission: {_join_limited(overview.get('failsafeEvents', []))}. "
        "These message and event streams can be cross-referenced directly in the app for deeper investigation."
    )
    return [paragraph_one, paragraph_two, paragraph_three]


def _build_system_health_paragraphs(analysis: dict[str, Any]) -> list[str]:
    overview = analysis["overview"]
    power = analysis["power"]
    vibration = analysis["vibration"]
    rc = analysis["rc"]
    paragraph_one = (
        f"Power analysis reported {_text(power.get('batteryStatus'))} with a minimum battery reading of {_text(power.get('minVoltage'))} and peak current of {_text(power.get('maxCurrent'))}. "
        f"Communication strength was {_text(overview.get('communicationStrength'))}, and signal strength was {_text(overview.get('signalStrength'))}."
    )
    paragraph_two = (
        f"Vibration review used {_text(vibration.get('sampleCount'), fallback='0')} samples. "
        f"The highest observed vibration was {_text(vibration.get('peakVibration'))}, supporting quick inspection of airframe and IMU stability."
    )
    paragraph_three = (
        f"RC system review identified {_text(overview.get('rcHealth'))} with up to {MAX_RC_CHANNELS} supported channels, and {_text(rc.get('activeChannelCount'), fallback='0')} channels were active in this mission."
    )
    return [paragraph_one, paragraph_two, paragraph_three]


def _build_findings_paragraphs(analysis: dict[str, Any]) -> list[str]:
    overview = analysis["overview"]
    reports = analysis["reports"]
    assessment = _assessment_label(overview, analysis["messages"])
    paragraph_one = (
        f"Overall mission assessment: {assessment}. "
        f"Recorded errors include {_join_limited(overview.get('errorMessages', []))}."
    )
    paragraph_two = (
        "The generated Excel workbook and PDF report are aligned into one export folder so the reviewed mission can be shared, archived, "
        "or attached to downstream operational reporting workflows without additional manual collation."
    )
    paragraph_three = (
        f"Report summary focus areas were {_join_limited(reports.get('highlights', []), fallback='mission overview, route review, system health, and event inspection')}."
    )
    return [paragraph_one, paragraph_two, paragraph_three]


def generate_excel_report(analysis: dict[str, Any], output_dir: str | None = None, report_folder: Path | None = None) -> str:
    report_folder = report_folder or _create_report_folder(output_dir)
    excel_path = report_folder / "flight_data.xlsx"

    overview = analysis["overview"]
    timeline = analysis["timeline"]
    messages = analysis["messages"]
    power = analysis["power"]
    vibration = analysis["vibration"]
    rc = analysis["rc"]
    map_data = analysis["map"]

    overview_df = pd.DataFrame(
        [
            ["Log Name", _text(overview.get("logName"))],
            ["Vehicle Type", _text(overview.get("vehicleType"))],
            ["Mission Date Time", _text(overview.get("dateTime"))],
            ["Flight Duration", _text(overview.get("totalFlightDuration"))],
            ["Mission Window", _text(overview.get("armDisarmTime"))],
            ["Flight Count", _text(overview.get("flightCount"), fallback="0")],
            ["Primary Modes", _text(overview.get("flightModes"))],
            ["GPS Status", _text(overview.get("gpsStatus"))],
            ["Satellite Count", _text(overview.get("satelliteCount"))],
            ["Home Location", _text(overview.get("homeLocation"))],
            ["Distance Traveled", _text(overview.get("distanceTraveled"))],
            ["Max Altitude", _text(overview.get("maxAltitude"))],
            ["Max Speed", _text(overview.get("maxSpeed"))],
            ["Orientation Source", _text(overview.get("orientationSource"))],
            ["IMU Count", _text(overview.get("imuCount"))],
            ["Proximity Sensors", _text(overview.get("proximitySensorCount"))],
            ["RC Health", _text(overview.get("rcHealth"))],
            ["Communication Strength", _text(overview.get("communicationStrength"))],
            ["Signal Strength", _text(overview.get("signalStrength"))],
            ["Failsafe Events", _text(overview.get("failsafeEvents"))],
            ["Warnings", _text(overview.get("keyWarnings"))],
            ["Anomalies", _text(overview.get("keyAnomalies"))],
            ["Errors", _text(overview.get("errorMessages"))],
        ],
        columns=["Metric", "Value"],
    )

    timeline_df = pd.DataFrame(timeline.get("events", []))
    if timeline_df.empty:
        timeline_df = pd.DataFrame(columns=["timeS", "title", "description", "severity"])

    message_records_df = pd.DataFrame(messages.get("records", []))
    if message_records_df.empty:
        message_records_df = pd.DataFrame(columns=["timeS", "source", "severity", "text"])

    message_counts_df = pd.DataFrame(messages.get("counts", []))
    if message_counts_df.empty:
        message_counts_df = pd.DataFrame(columns=["label", "count"])

    power_df = pd.DataFrame(power.get("samples", []))
    if power_df.empty:
        power_df = pd.DataFrame(columns=["timeS", "voltage", "current", "remainingPct"])

    vibration_df = pd.DataFrame(vibration.get("samples", []))
    if vibration_df.empty:
        vibration_df = pd.DataFrame(columns=["timeS", "x", "y", "z", "magnitude"])

    map_route_df = pd.DataFrame(map_data.get("routePoints", []))
    if map_route_df.empty:
        map_route_df = pd.DataFrame(columns=["timeS", "lat", "lon", "alt", "speed"])

    rc_df = pd.DataFrame(rc.get("samples", []))
    if rc_df.empty:
        rc_df = pd.DataFrame(columns=["timeS", *RC_CHANNEL_COLUMNS])

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        overview_df.to_excel(writer, sheet_name="Overview", index=False)
        timeline_df.to_excel(writer, sheet_name="Timeline", index=False)
        message_records_df.to_excel(writer, sheet_name="Messages", index=False)
        message_counts_df.to_excel(writer, sheet_name="MessageCounts", index=False)
        power_df.to_excel(writer, sheet_name="Power", index=False)
        vibration_df.to_excel(writer, sheet_name="Vibration", index=False)
        rc_df.to_excel(writer, sheet_name="RC", index=False)
        map_route_df.to_excel(writer, sheet_name="MapRoute", index=False)

        workbook = writer.book
        power_sheet = writer.sheets["Power"]
        vibration_sheet = writer.sheets["Vibration"]
        rc_sheet = writer.sheets["RC"]
        map_route_sheet = writer.sheets["MapRoute"]
        message_counts_sheet = writer.sheets["MessageCounts"]

        _add_line_chart(
            power_sheet,
            title="Battery Voltage And Current",
            y_axis_title="Voltage / Current",
            x_axis_title="Time",
            min_col=2,
            max_col=3,
            header_row=1,
            first_data_row=2,
            last_data_row=len(power_df) + 1,
            anchor="F2",
            colors=["FF4A4A", "0F8BFF"],
        )
        _add_line_chart(
            vibration_sheet,
            title="Vibration Magnitude",
            y_axis_title="Magnitude",
            x_axis_title="Time",
            min_col=5,
            max_col=5,
            header_row=1,
            first_data_row=2,
            last_data_row=len(vibration_df) + 1,
            anchor="G2",
            colors=["F97316"],
        )
        active_rc_columns = [column for column in RC_CHANNEL_COLUMNS if column in rc_df.columns and rc_df[column].notna().any()]
        if active_rc_columns:
            first_col = rc_df.columns.get_loc(active_rc_columns[0]) + 1
            last_col = rc_df.columns.get_loc(active_rc_columns[-1]) + 1
            _add_line_chart(
                rc_sheet,
                title="RC Channel Trends",
                y_axis_title="PWM",
                x_axis_title="Time",
                min_col=first_col,
                max_col=last_col,
                header_row=1,
                first_data_row=2,
                last_data_row=len(rc_df) + 1,
                anchor="R2",
                colors=["FF4A4A", "F97316", "F59E0B", "84CC16", "14B8A6", "0F8BFF", "6366F1", "EC4899"],
                width=22,
                height=11,
            )
        _add_line_chart(
            map_route_sheet,
            title="Route Altitude and Speed",
            y_axis_title="Altitude / Speed",
            x_axis_title="Time",
            min_col=4,
            max_col=5,
            header_row=1,
            first_data_row=2,
            last_data_row=len(map_route_df) + 1,
            anchor="J2",
            colors=["FF4A4A", "0F8BFF"],
        )
        _add_bar_chart(
            message_counts_sheet,
            title="Message Counts",
            y_axis_title="Count",
            min_col=2,
            max_col=2,
            header_row=1,
            first_data_row=2,
            last_data_row=len(message_counts_df) + 1,
            anchor="E2",
            colors=["FF4A4A"],
        )

        for sheet_name in workbook.sheetnames:
            _autosize_columns(workbook[sheet_name])

    return str(excel_path)


def generate_pdf_report(analysis: dict[str, Any], output_dir: str | None = None, report_folder: Path | None = None) -> str:
    report_folder = report_folder or _create_report_folder(output_dir)
    pdf_path = report_folder / "mission_report.pdf"
    generated_at = datetime.now()
    current_year = generated_at.year
    overview = analysis["overview"]
    reports = analysis["reports"]

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="KapYahHeaderTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=28,
            textColor=colors.HexColor("#D62828"),
            alignment=TA_LEFT,
            spaceAfter=3,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KapYahHeaderSub",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=14,
            textColor=colors.HexColor("#425166"),
            alignment=TA_LEFT,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KapYahSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#18202B"),
            spaceBefore=10,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KapYahBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=10,
            leading=15,
            textColor=colors.HexColor("#263444"),
            spaceAfter=7,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KapYahFooter",
            parent=styles["BodyText"],
            alignment=TA_CENTER,
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#6E7A8B"),
            spaceBefore=12,
        )
    )

    story: list[Any] = []
    logo_path = _resolve_logo_path()
    title_block = [
        Paragraph("KapYah Mission Analysis Report", styles["KapYahHeaderTitle"]),
    ]
    if logo_path and logo_path.exists():
        logo = Image(str(logo_path), width=18 * mm, height=18 * mm)
        header_table = Table([[logo, title_block]], colWidths=[24 * mm, 146 * mm])
    else:
        header_table = Table([[title_block]], colWidths=[170 * mm])
    header_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (0, 0), "LEFT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    story.append(header_table)
    story.append(Spacer(1, 5 * mm))

    assessment = _assessment_label(overview, analysis["messages"])
    meta_data = [
        ["Generated", generated_at.strftime("%d %b %Y, %I:%M %p")],
        ["Website", "www.kapyah.com"],
        ["Log File", _text(overview.get("logName"))],
        ["Vehicle", _text(overview.get("vehicleType"))],
        ["Mission Status", assessment],
    ]
    meta_table = Table(meta_data, colWidths=[34 * mm, 126 * mm])
    meta_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.2),
                ("LEADING", (0, 0), (-1, -1), 12),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#D62828")),
                ("TEXTCOLOR", (1, 0), (1, -1), colors.HexColor("#263444")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, -1), (-1, -1), 0.6, colors.HexColor("#E3CFC5")),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    story.append(meta_table)
    story.append(Spacer(1, 4 * mm))

    glance_data = [
        ["Flight Duration", _text(overview.get("totalFlightDuration"))],
        ["Mission Window", _text(overview.get("armDisarmTime"))],
        ["GPS Status", _text(overview.get("gpsStatus"))],
        ["Satellite Count", _text(overview.get("satelliteCount"))],
        ["Home Location Details", _text(overview.get("homeLocation"))],
        ["Distance Traveled", _text(overview.get("distanceTraveled"))],
        ["Max Height Reached", _text(overview.get("maxAltitude"))],
        ["Max Speed", _text(overview.get("maxSpeed"))],
        ["Orientation Source", _text(overview.get("orientationSource"))],
        ["IMU Summary", _text(overview.get("imuCount"))],
        ["Proximity Sensors", _text(overview.get("proximitySensorCount"))],
        ["RC Health", _text(overview.get("rcHealth"))],
    ]
    story.append(Paragraph("Mission At A Glance", styles["KapYahSection"]))
    glance_table = Table(glance_data, colWidths=[52 * mm, 108 * mm])
    glance_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.2),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#18202B")),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.HexColor("#FFF8F4"), colors.whitesmoke]),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#E3CFC5")),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(glance_table)
    story.append(Spacer(1, 3 * mm))

    sections = [
        ("Mission Overview", _build_mission_overview_paragraphs(analysis, generated_at)),
        ("Route And GPS Review", _build_navigation_paragraphs(analysis)),
        ("Timeline And Messages", _build_timeline_message_paragraphs(analysis)),
        ("Power, Vibration, And RC Review", _build_system_health_paragraphs(analysis)),
        ("Findings And Final Assessment", _build_findings_paragraphs(analysis)),
    ]
    for title, paragraphs in sections:
        story.append(Paragraph(title, styles["KapYahSection"]))
        for paragraph in paragraphs:
            story.append(Paragraph(escape(paragraph), styles["KapYahBody"]))

    story.append(Paragraph(f"&#169; {current_year} KapYah Industries Pvt. Ltd. All rights reserved.", styles["KapYahFooter"]))

    document = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=14 * mm,
    )
    document.build(story)
    return str(pdf_path)
